import hashlib
import re
from datetime import date
from io import BytesIO
from pathlib import Path

import fitz
from docx import Document

from app.models.documents import (
    CandidateExtraction,
    DocumentPage,
    Evidence,
    ExtractedField,
    ExtractedSkill,
    ParsedDocument,
)
from app.models.mvp import JobParseRequest, JobParseResult, JobRequirement, ResumeAnalysis

PARSER_VERSION = "local-parser-v1"
EXTRACTION_VERSION = "local-rules-v1"

SKILL_ALIASES: dict[str, tuple[str, ...]] = {
    "Go": ("golang",),
    "Python": (),
    "Java": (),
    "JavaScript": ("javascript", "js"),
    "TypeScript": ("typescript", "ts"),
    "React": ("react.js", "reactjs"),
    "Next.js": ("nextjs", "next js"),
    "Node.js": ("nodejs", "node js"),
    "C#": ("csharp",),
    ".NET": ("dotnet", "asp.net"),
    "PHP": (),
    "Ruby": (),
    "Rust": (),
    "C++": ("cpp",),
    "AWS": ("amazon web services",),
    "Azure": (),
    "GCP": ("google cloud", "google cloud platform"),
    "Docker": (),
    "Kubernetes": ("k8s", "aws eks", "eks"),
    "PostgreSQL": ("postgres", "postgresql"),
    "MySQL": (),
    "SQL": (),
    "MongoDB": ("mongo",),
    "Redis": (),
    "Kafka": ("apache kafka",),
    "RabbitMQ": (),
    "GraphQL": (),
    "REST": ("rest api", "restful"),
    "gRPC": (),
    "Microservices": ("microservice", "micro-services"),
    "Linux": (),
    "Git": (),
    "Terraform": (),
    "CI/CD": ("cicd", "continuous integration", "continuous delivery"),
    "Machine Learning": ("machine learning", "ml"),
    "NLP": ("natural language processing",),
    "PyTorch": (),
    "TensorFlow": (),
    "FastAPI": (),
    "Django": (),
    "Flask": (),
    "Spring Boot": ("springboot",),
    "Gin": (),
    "Chi": (),
    "HTML": (),
    "CSS": (),
}

CONTACT_LINE = re.compile(
    r"(?i)(?:[\w.+-]+@[\w.-]+\.[a-z]{2,}|\+?\d[\d\s().-]{7,}\d|linkedin\.com)"
)
YEAR_PATTERN = re.compile(r"(?i)(\d{1,2}(?:\.\d+)?)\s*\+?\s*(?:years?|yrs?|năm)")
DATE_RANGE_PATTERN = re.compile(
    r"(?i)(?:\d{1,2}[/-])?(\d{4})\s*(?:-|–|—|to|đến)\s*"
    r"((?:\d{1,2}[/-])?\d{4}|present|hiện tại|now|nay)\b"
)
TITLE_PATTERN = re.compile(
    r"(?i)\b(engineer|developer|architect|scientist|manager|specialist|consultant|"
    r"kỹ sư|lập trình viên|trưởng nhóm|quản lý)\b"
)

DEGREE_PATTERN = re.compile(
    r"(?i)\b(tiến\s*sĩ|ph\.?d\.?|thạc\s*sĩ|master(?:'s)?\s*degree|master(?:'s)?|msc|mba|"
    r"cử\s*nhân|kỹ\s*sư|bachelor(?:'s)?\s*degree|bachelor(?:'s)?|bsc|cao\s*đẳng|"
    r"associate\s*degree)\b"
)
INSTITUTION_PATTERN = re.compile(
    r"(?i)\b(?:trường\s+)?(đại học|học viện|university|college|institute)\b[^\n,;]{0,80}"
)
FIELD_OF_STUDY_PATTERN = re.compile(
    r"(?i)(?:chuyên ngành|ngành học|ngành|major(?:ed)?\s+in|field of study)[:\s]+"
    r"([^\n,;.]{2,80})"
)
CERTIFICATION_PATTERN = re.compile(
    r"(?i)\b(chứng chỉ\s+[^\n,;.]{2,80}|certified\s+[^\n,;.]{2,80}|"
    r"certificate\s+(?:in|of)\s+[^\n,;.]{2,80}|aws certified\s+[^\n,;.]{0,60}|"
    r"pmp|scrum master(?:\s*certified)?|comptia\s+[^\n,;.]{0,40}|ccna|cissp|"
    r"toeic\s*[\d.]*|ielts\s*[\d.]*)\b"
)
LANGUAGE_NAMES: dict[str, str] = {
    "tiếng anh": "English",
    "english": "English",
    "tiếng nhật": "Japanese",
    "japanese": "Japanese",
    "tiếng trung": "Chinese",
    "mandarin": "Chinese",
    "tiếng hàn": "Korean",
    "korean": "Korean",
    "tiếng pháp": "French",
    "french": "French",
    "tiếng đức": "German",
    "german": "German",
    "tiếng việt": "Vietnamese",
    "vietnamese": "Vietnamese",
}
PROFICIENCY_PATTERN = re.compile(
    r"(?i)\b(native|fluent|advanced|professional working proficiency|intermediate|"
    r"elementary|basic|proficient|thành thạo|trôi chảy|khá tốt|khá|cơ bản|"
    r"giao tiếp tốt|giao tiếp)\b"
)


def _phrase_pattern(phrase: str) -> re.Pattern[str]:
    if phrase == "Go":
        return re.compile(r"\b(?:Go|Golang)\b")
    if phrase in {"C++", "C#", ".NET", "CI/CD"}:
        return re.compile(re.escape(phrase), re.IGNORECASE)
    return re.compile(r"(?<![\w])" + re.escape(phrase) + r"(?![\w])", re.IGNORECASE)


def normalize_skill(value: str) -> str | None:
    for canonical, aliases in SKILL_ALIASES.items():
        if value.casefold() in {canonical.casefold(), *(alias.casefold() for alias in aliases)}:
            return canonical
    return None


def _skill_mentions(text: str) -> list[tuple[str, re.Match[str]]]:
    found: list[tuple[str, re.Match[str]]] = []
    for canonical, aliases in SKILL_ALIASES.items():
        matches: list[re.Match[str]] = []
        for phrase in (canonical, *aliases):
            matches.extend(_phrase_pattern(phrase).finditer(text))
        if matches:
            found.append((canonical, min(matches, key=lambda match: match.start())))
    return sorted(found, key=lambda item: item[1].start())


def _line_excerpt(text: str, offset: int, limit: int = 500) -> tuple[str, int, int]:
    start = text.rfind("\n", 0, offset) + 1
    end = text.find("\n", offset)
    if end < 0:
        end = len(text)
    excerpt = text[start:end].strip()
    if not excerpt:
        start = max(0, offset - 100)
        end = min(len(text), offset + 200)
        excerpt = text[start:end].strip()
    return excerpt[:limit], start, min(end, start + limit)


def _page_for_offset(document: ParsedDocument, offset: int) -> int | None:
    for page in document.pages:
        if page.start_offset <= offset <= page.end_offset:
            return page.number
    return None


def _evidence(document: ParsedDocument, text: str, offset: int, prefix: str) -> Evidence:
    excerpt, start, end = _line_excerpt(text, offset)
    digest = hashlib.sha1(f"{prefix}:{start}:{excerpt}".encode()).hexdigest()[:20]
    return Evidence(
        id=f"{prefix}-{digest}",
        text=excerpt,
        page=_page_for_offset(document, start),
        start_offset=start,
        end_offset=end,
    )


def parse_document(filename: str, content_type: str, data: bytes) -> ParsedDocument:
    suffix = Path(filename).suffix.lower()
    if content_type == "application/pdf" or suffix == ".pdf":
        source = fitz.open(stream=data, filetype="pdf")
        raw_pages = [page.get_text("text").strip() for page in source]
        source.close()
    elif (
        content_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        or suffix == ".docx"
    ):
        document = Document(BytesIO(data))
        raw_pages = ["\n".join(paragraph.text for paragraph in document.paragraphs).strip()]
    else:
        raise ValueError("only PDF and DOCX files are supported")

    if not any(raw_pages):
        raise ValueError("the document does not contain an extractable text layer")
    pages: list[DocumentPage] = []
    parts: list[str] = []
    offset = 0
    for number, page_text in enumerate(raw_pages, start=1):
        if parts:
            parts.append("\n")
            offset += 1
        start = offset
        parts.append(page_text)
        offset += len(page_text)
        pages.append(
            DocumentPage(
                number=number,
                text=page_text,
                start_offset=start,
                end_offset=offset,
            )
        )
    return ParsedDocument(
        text="".join(parts),
        pages=pages,
        metadata={"filename": filename, "page_count": len(pages)},
        parser_version=PARSER_VERSION,
        used_ocr=False,
    )


def _skill_estimated_years(text: str, canonical: str) -> float | None:
    years_found: list[float] = []
    for phrase in (canonical, *SKILL_ALIASES[canonical]):
        for match in _phrase_pattern(phrase).finditer(text):
            context_start = max(0, match.start() - 80)
            context_end = min(len(text), match.end() + 80)
            years_match = YEAR_PATTERN.search(text[context_start:context_end])
            if years_match:
                years_found.append(float(years_match.group(1)))
    return max(years_found) if years_found else None


def _estimate_years_from_date_ranges(text: str) -> tuple[float, int] | None:
    current_year = date.today().year
    ranges: list[tuple[int, int]] = []
    first_offset: int | None = None
    for match in DATE_RANGE_PATTERN.finditer(text):
        start_year = int(match.group(1))
        end_token = match.group(2).casefold()
        if end_token in {"present", "hiện tại", "now", "nay"}:
            end_year = current_year
        else:
            year_digits = re.search(r"\d{4}", end_token)
            if not year_digits:
                continue
            end_year = int(year_digits.group(0))
        if end_year < start_year or start_year < 1970 or end_year > current_year + 1:
            continue
        ranges.append((start_year, end_year))
        if first_offset is None:
            first_offset = match.start()
    if not ranges or first_offset is None:
        return None
    ranges.sort()
    merged: list[list[int]] = []
    for start_year, end_year in ranges:
        if merged and start_year <= merged[-1][1]:
            merged[-1][1] = max(merged[-1][1], end_year)
        else:
            merged.append([start_year, end_year])
    total_years = sum(end_year - start_year for start_year, end_year in merged)
    return (float(total_years), first_offset) if total_years > 0 else None


def _extract_educations(document: ParsedDocument, text: str) -> list[dict]:
    educations: list[dict] = []
    seen: set[str] = set()
    for match in INSTITUTION_PATTERN.finditer(text):
        excerpt, start, _ = _line_excerpt(text, match.start())
        key = excerpt.casefold()
        if not excerpt or key in seen:
            continue
        seen.add(key)
        degree_match = DEGREE_PATTERN.search(excerpt)
        field_match = FIELD_OF_STUDY_PATTERN.search(excerpt)
        evidence = _evidence(document, text, start, "education")
        educations.append(
            {
                "institution": excerpt[:200],
                "degree": degree_match.group(0).strip() if degree_match else "",
                "field_of_study": field_match.group(1).strip() if field_match else "",
                "confidence": 0.75 if degree_match else 0.6,
                "evidences": [evidence.model_dump(mode="json")],
            }
        )
    for match in DEGREE_PATTERN.finditer(text):
        excerpt, start, _ = _line_excerpt(text, match.start())
        key = excerpt.casefold()
        if not excerpt or key in seen:
            continue
        seen.add(key)
        field_match = FIELD_OF_STUDY_PATTERN.search(excerpt)
        evidence = _evidence(document, text, start, "education")
        educations.append(
            {
                "institution": "",
                "degree": match.group(0).strip(),
                "field_of_study": field_match.group(1).strip() if field_match else "",
                "confidence": 0.55,
                "evidences": [evidence.model_dump(mode="json")],
            }
        )
    return educations[:10]


def _extract_certifications(document: ParsedDocument, text: str) -> list[dict]:
    certifications: list[dict] = []
    seen: set[str] = set()
    for match in CERTIFICATION_PATTERN.finditer(text):
        name = re.sub(r"\s+", " ", match.group(0)).strip()
        key = name.casefold()
        if not name or key in seen:
            continue
        seen.add(key)
        evidence = _evidence(document, text, match.start(), "certification")
        certifications.append(
            {
                "name": name[:200],
                "issuer": "",
                "confidence": 0.7,
                "evidences": [evidence.model_dump(mode="json")],
            }
        )
    return certifications[:10]


def _extract_languages(document: ParsedDocument, text: str) -> list[dict]:
    languages: list[dict] = []
    seen: set[str] = set()
    for phrase, canonical in LANGUAGE_NAMES.items():
        if canonical in seen:
            continue
        match = re.search(r"(?<![\w])" + re.escape(phrase) + r"(?![\w])", text, re.IGNORECASE)
        if not match:
            continue
        seen.add(canonical)
        context_start = max(0, match.start() - 60)
        context_end = min(len(text), match.end() + 60)
        context = text[context_start:context_end]
        proficiency_match = PROFICIENCY_PATTERN.search(context)
        proficiency = proficiency_match.group(0).strip() if proficiency_match else ""
        if not proficiency:
            score_match = re.search(r"(?i)(ielts\s*[\d.]+|toeic\s*\d+)", context)
            proficiency = score_match.group(0).strip() if score_match else ""
        evidence = _evidence(document, text, match.start(), f"language-{canonical.casefold()}")
        languages.append(
            {
                "language": canonical,
                "proficiency": proficiency,
                "confidence": 0.8 if proficiency else 0.6,
                "evidences": [evidence.model_dump(mode="json")],
            }
        )
    return languages


def _field(value: str | float, evidence: Evidence) -> ExtractedField:
    return ExtractedField(
        value=value,
        confidence=0.88,
        evidences=[evidence],
        extraction_method="rule",
        model_version=EXTRACTION_VERSION,
    )


def analyze_resume(filename: str, content_type: str, data: bytes) -> ResumeAnalysis:
    document = parse_document(filename, content_type, data)
    text = document.text
    lines = [(line.strip(), text.find(line)) for line in text.splitlines() if line.strip()]
    email_match = re.search(r"(?i)[\w.+-]+@[\w.-]+\.[a-z]{2,}", text)
    phone_match = re.search(r"(?<!\w)\+?\d[\d\s().-]{7,}\d(?!\w)", text)

    name_line = next(
        (
            (line, offset)
            for line, offset in lines[:8]
            if len(line) <= 100 and not CONTACT_LINE.search(line) and not TITLE_PATTERN.search(line)
        ),
        None,
    )
    title_line = next(
        ((line, offset) for line, offset in lines[:20] if TITLE_PATTERN.search(line)),
        None,
    )
    years = [float(match.group(1)) for match in YEAR_PATTERN.finditer(text)]
    year_pattern_match = next(YEAR_PATTERN.finditer(text), None)
    inferred_experience = _estimate_years_from_date_ranges(text)

    total_years_value: float | None = max(years) if years else None
    total_years_offset: int | None = year_pattern_match.start() if year_pattern_match else None
    if inferred_experience is not None:
        inferred_years, inferred_offset = inferred_experience
        if total_years_value is None or inferred_years > total_years_value:
            total_years_value = inferred_years
            total_years_offset = inferred_offset

    skills: list[ExtractedSkill] = []
    for canonical, match in _skill_mentions(text):
        evidence = _evidence(document, text, match.start(), f"skill-{canonical.casefold()}")
        skills.append(
            ExtractedSkill(
                name=match.group(0),
                normalized_name=canonical,
                match_kind="explicit",
                estimated_years=_skill_estimated_years(text, canonical),
                confidence=0.95,
                evidences=[evidence],
            )
        )

    safe_summary_lines = [
        line
        for line, _ in lines[:12]
        if not CONTACT_LINE.search(line) and line != (name_line or ("", 0))[0]
    ]
    summary_text = " · ".join(safe_summary_lines)[:600] or "CV đã được trích xuất."
    summary_offset = text.find(safe_summary_lines[0]) if safe_summary_lines else 0
    summary_evidence = _evidence(document, text, max(0, summary_offset), "summary")

    return ResumeAnalysis(
        document=document,
        candidate=CandidateExtraction(
            full_name=(
                _field(name_line[0], _evidence(document, text, name_line[1], "name"))
                if name_line
                else None
            ),
            email=(
                _field(
                    email_match.group(0),
                    _evidence(document, text, email_match.start(), "email"),
                )
                if email_match
                else None
            ),
            phone=(
                _field(
                    phone_match.group(0).strip(),
                    _evidence(document, text, phone_match.start(), "phone"),
                )
                if phone_match
                else None
            ),
            current_title=(
                _field(title_line[0], _evidence(document, text, title_line[1], "title"))
                if title_line
                else None
            ),
            summary=_field(summary_text, summary_evidence),
            total_years_experience=(
                _field(
                    total_years_value,
                    _evidence(document, text, total_years_offset or 0, "experience"),
                )
                if total_years_value is not None
                else None
            ),
            skills=skills,
            educations=_extract_educations(document, text),
            certifications=_extract_certifications(document, text),
            languages=_extract_languages(document, text),
            extraction_model_version=EXTRACTION_VERSION,
        ),
    )


def parse_job(payload: JobParseRequest) -> JobParseResult:
    description = payload.description
    requirements: list[JobRequirement] = []
    counter = 1
    for canonical, match in _skill_mentions(description):
        excerpt, start, end = _line_excerpt(description, match.start())
        lowered = excerpt.casefold()
        nice = any(
            marker in lowered
            for marker in ("nice to have", "preferred", "plus", "ưu tiên", "lợi thế")
        )
        years_match = YEAR_PATTERN.search(excerpt)
        requirements.append(
            JobRequirement(
                requirement_key=f"R{counter}",
                name=canonical,
                requirement_type="skill",
                priority="nice_to_have" if nice else "must_have",
                weight=0.6 if nice else 1.0,
                minimum_years=float(years_match.group(1)) if years_match else None,
                is_hard_constraint=False,
                description=excerpt,
                source_evidence=[
                    Evidence(
                        id=f"jd-{counter}",
                        text=excerpt,
                        start_offset=start,
                        end_offset=end,
                    )
                ],
            )
        )
        counter += 1

    general_years = YEAR_PATTERN.search(description)
    if general_years:
        excerpt, start, end = _line_excerpt(description, general_years.start())
        requirements.append(
            JobRequirement(
                requirement_key=f"R{counter}",
                name="Kinh nghiệm liên quan",
                requirement_type="experience",
                priority="must_have",
                weight=1.0,
                minimum_years=float(general_years.group(1)),
                description=excerpt,
                source_evidence=[
                    Evidence(
                        id=f"jd-{counter}",
                        text=excerpt,
                        start_offset=start,
                        end_offset=end,
                    )
                ],
            )
        )
        counter += 1

    degree_match = DEGREE_PATTERN.search(description)
    if degree_match:
        excerpt, start, end = _line_excerpt(description, degree_match.start())
        lowered = excerpt.casefold()
        nice = any(
            marker in lowered
            for marker in ("nice to have", "preferred", "plus", "ưu tiên", "lợi thế")
        )
        requirements.append(
            JobRequirement(
                requirement_key=f"R{counter}",
                name=f"Học vấn: {degree_match.group(0).strip()}",
                requirement_type="education",
                priority="nice_to_have" if nice else "must_have",
                weight=0.6 if nice else 0.8,
                description=excerpt,
                source_evidence=[
                    Evidence(id=f"jd-{counter}", text=excerpt, start_offset=start, end_offset=end)
                ],
            )
        )
        counter += 1

    cert_match = CERTIFICATION_PATTERN.search(description)
    if cert_match:
        excerpt, start, end = _line_excerpt(description, cert_match.start())
        lowered = excerpt.casefold()
        nice = any(
            marker in lowered
            for marker in ("nice to have", "preferred", "plus", "ưu tiên", "lợi thế")
        )
        requirements.append(
            JobRequirement(
                requirement_key=f"R{counter}",
                name=f"Chứng chỉ: {re.sub(r'\\s+', ' ', cert_match.group(0)).strip()}",
                requirement_type="certification",
                priority="nice_to_have" if nice else "must_have",
                weight=0.5 if nice else 0.7,
                description=excerpt,
                source_evidence=[
                    Evidence(id=f"jd-{counter}", text=excerpt, start_offset=start, end_offset=end)
                ],
            )
        )
        counter += 1

    for phrase, canonical in LANGUAGE_NAMES.items():
        if canonical == "Vietnamese":
            continue
        language_match = re.search(
            r"(?<![\w])" + re.escape(phrase) + r"(?![\w])", description, re.IGNORECASE
        )
        if not language_match:
            continue
        excerpt, start, end = _line_excerpt(description, language_match.start())
        lowered = excerpt.casefold()
        nice = any(
            marker in lowered
            for marker in ("nice to have", "preferred", "plus", "ưu tiên", "lợi thế")
        )
        requirements.append(
            JobRequirement(
                requirement_key=f"R{counter}",
                name=f"Ngoại ngữ: {canonical}",
                requirement_type="language",
                priority="nice_to_have" if nice else "must_have",
                weight=0.5 if nice else 0.7,
                description=excerpt,
                source_evidence=[
                    Evidence(id=f"jd-{counter}", text=excerpt, start_offset=start, end_offset=end)
                ],
            )
        )
        counter += 1
        break

    requirements.append(
        JobRequirement(
            requirement_key=f"R{counter}",
            name=payload.title,
            requirement_type="other",
            priority="nice_to_have",
            weight=0.5,
            description="Mức độ tương đồng giữa chức danh hiện tại và vị trí tuyển dụng.",
        )
    )
    seniority = next(
        (
            value
            for value in ("principal", "lead", "senior", "mid", "junior", "intern")
            if value in f"{payload.title} {description[:500]}".casefold()
        ),
        None,
    )
    summary = re.sub(r"\s+", " ", description).strip()[:800]
    return JobParseResult(
        title=payload.title,
        summary=summary,
        seniority=seniority,
        requirements=requirements,
        parser_version="jd-rules-v1",
    )
